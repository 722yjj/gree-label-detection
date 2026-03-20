import os
import sys

from openpyxl import load_workbook
from pydantic import BaseModel

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from label_detection.workflows.unified import compare_text_results


class DummyLabel(BaseModel):
    manufacturer: str | None = None
    weight: str | None = None


def test_compare_text_results_uses_normalized_field_matching(tmp_path):
    template = DummyLabel(
        manufacturer="GREE ELECTRIC APPLIANCES,INC.OF ZHUHAI",
        weight="13.5kg",
    )
    target = DummyLabel(
        manufacturer="GREE ELECTRIC APPLIANCES, INC. OF ZHUHAI",
        weight="135kg",
    )

    output_path = tmp_path / "text_comparison.xlsx"
    compare_text_results(template, target, str(output_path))

    workbook = load_workbook(output_path)
    worksheet = workbook["文字对比"]

    assert worksheet.cell(row=2, column=4).value == "✓"
    assert worksheet.cell(row=3, column=4).value == "✗"
