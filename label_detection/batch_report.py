"""Generate Excel / PDF reports with template / target / diff thumbnails per case."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from label_detection.core.config import PROJECT_ROOT

# Layout constants – Excel
_THUMB_MAX_HEIGHT = 200  # px – thumbnail fits in ~200 px tall row
_ROW_HEIGHT_PT = 155  # openpyxl row height in points (≈ 200 px)
_COL_WIDTH_CHAR = 38  # column width in character-units (≈ 270 px)

REPORT_FILENAME = "batch_report.xlsx"
PDF_REPORT_FILENAME = "batch_report.pdf"


def _resolve_path(raw: str) -> Optional[Path]:
    """Resolve a path string that may be repo-relative or absolute."""
    p = Path(raw)
    if p.is_absolute() and p.exists():
        return p
    candidate = PROJECT_ROOT / raw
    if candidate.exists():
        return candidate
    return None


def _add_thumbnail(ws, image_path: Optional[Path], cell_ref: str) -> None:
    """Embed a resized image into *cell_ref* if *image_path* exists."""
    if image_path is None or not image_path.exists():
        ws[cell_ref] = "(无图)"
        return
    try:
        img = XlImage(str(image_path))
        # Scale to fit _THUMB_MAX_HEIGHT while keeping aspect ratio
        if img.height and img.height > _THUMB_MAX_HEIGHT:
            ratio = _THUMB_MAX_HEIGHT / img.height
            img.width = int(img.width * ratio)
            img.height = _THUMB_MAX_HEIGHT
        img.anchor = cell_ref
        ws.add_image(img)
    except Exception:
        ws[cell_ref] = "(图片加载失败)"


def generate_batch_excel_report(
    output_dir: Path,
    records: Sequence,
    summary_payload: Optional[Dict] = None,
) -> Path:
    """Build an Excel workbook with one row per case (template / target / diff).

    Parameters
    ----------
    output_dir:
        Batch output root – the report is saved here.
    records:
        Sequence of ``CaseRunRecord`` dataclass instances.
    summary_payload:
        Optional overall summary dict (unused for now, reserved for a
        summary sheet).

    Returns
    -------
    Path to the generated ``.xlsx`` file.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "对比报告"

    # ── Header ──
    headers = ["编号", "Case ID", "模板图", "实拍图", "差异图", "判定"]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)

    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Set column widths
    ws.column_dimensions["A"].width = 18  # 编号
    ws.column_dimensions["B"].width = 32  # Case ID
    for col_letter in ("C", "D", "E"):  # image columns
        ws.column_dimensions[col_letter].width = _COL_WIDTH_CHAR
    ws.column_dimensions["F"].width = 40  # 判定

    ws.row_dimensions[1].height = 25

    # ── Data rows ──
    for row_offset, record in enumerate(records):
        row = row_offset + 2  # 1-indexed, header is row 1
        ws.row_dimensions[row].height = _ROW_HEIGHT_PT

        # A: 编号
        ws.cell(row=row, column=1, value=record.code).alignment = Alignment(
            vertical="center", horizontal="center"
        )

        # B: Case ID
        ws.cell(row=row, column=2, value=record.case_id).alignment = Alignment(
            vertical="center", wrap_text=True
        )

        # Load result.json to get original template / target paths
        result_json_path = _resolve_path(record.result_json)
        result_data: Dict = {}
        if result_json_path is not None and result_json_path.exists():
            try:
                result_data = json.loads(result_json_path.read_text(encoding="utf-8"))
            except (json.JSONDecodeError, OSError):
                pass

        # C: 模板图 – prefer preprocessed jpg in case dir over raw PDF
        template_path = _resolve_path(record.template_path)
        case_dir = result_json_path.parent if result_json_path else None
        preprocessed_template = case_dir / "template_preprocessed.jpg" if case_dir else None
        if preprocessed_template and preprocessed_template.exists():
            template_path = preprocessed_template
        _add_thumbnail(ws, template_path, f"C{row}")

        # D: 实拍图 – prefer preprocessed jpg in case dir
        target_path = _resolve_path(record.target_path)
        preprocessed_target = case_dir / "target_preprocessed.jpg" if case_dir else None
        if preprocessed_target and preprocessed_target.exists():
            target_path = preprocessed_target
        _add_thumbnail(ws, target_path, f"D{row}")

        # E: 差异图
        vis_path = _resolve_path(record.visualization_diff) if record.visualization_diff else None
        _add_thumbnail(ws, vis_path, f"E{row}")

        # F: 判定
        verdict = record.verdict or result_data.get("verdict") or "(未知)"
        cell_verdict = ws.cell(row=row, column=6, value=verdict)
        cell_verdict.alignment = Alignment(vertical="center", wrap_text=True)
        # Colour by verdict
        if "❌" in verdict or "差异" in verdict:
            cell_verdict.font = Font(color="FF0000", bold=True)
        elif "⚠" in verdict:
            cell_verdict.font = Font(color="FF8C00", bold=True)
        elif "✅" in verdict:
            cell_verdict.font = Font(color="008000", bold=True)

    report_path = Path(output_dir) / REPORT_FILENAME
    wb.save(str(report_path))
    return report_path


# ---------------------------------------------------------------------------
# PDF report (fpdf2)
# ---------------------------------------------------------------------------

def _resolve_image_for_record(record, result_json_path: Optional[Path], role: str) -> Optional[Path]:
    """Find the best available image for *role* ('template' | 'target' | 'diff')."""
    case_dir = result_json_path.parent if result_json_path else None

    if role == "diff":
        if record.visualization_diff:
            return _resolve_path(record.visualization_diff)
        return None

    raw_path = record.template_path if role == "template" else record.target_path
    preprocessed_name = "template_preprocessed.jpg" if role == "template" else "target_preprocessed.jpg"

    # Prefer preprocessed image (handles PDF templates)
    if case_dir:
        pp = case_dir / preprocessed_name
        if pp.exists():
            return pp

    resolved = _resolve_path(raw_path) if raw_path else None
    if resolved and resolved.suffix.lower() not in (".pdf",):
        return resolved
    return None


def generate_batch_pdf_report(
    output_dir: Path,
    records: Sequence,
    summary_payload: Optional[Dict] = None,
) -> Path:
    """Build a PDF with one page per case showing template / target / diff images.

    Returns the path to the generated PDF file.
    """
    import re
    from fpdf import FPDF

    def _strip_emoji(text: str) -> str:
        """Remove emoji and other non-BMP characters that even CJK fonts lack."""
        return re.sub(
            r"[\U0001F300-\U0001FAFF\u2705\u274C\u26A0\uFE0F\U0001F4DD\U0001F3F7\U0001F5BC]",
            "",
            text,
        ).strip()

    # Find a CJK font on the system – prefer full-charset fonts (ASCII + CJK)
    _CJK_FONT_CANDIDATES = [
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/usr/share/fonts/wqy-zenhei/wqy-zenhei.ttc",
        "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf",
    ]
    cjk_font_path = None
    for candidate in _CJK_FONT_CANDIDATES:
        if Path(candidate).exists():
            cjk_font_path = candidate
            break

    # Landscape A4: 297 x 210 mm
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=False)

    if cjk_font_path:
        pdf.add_font("CJK", "", cjk_font_path)
        pdf.add_font("CJK", "B", cjk_font_path)
        font_name = "CJK"
    else:
        font_name = "Helvetica"

    page_w, page_h = 297, 210
    margin = 10
    usable_w = page_w - 2 * margin
    img_gap = 5  # gap between images
    img_w = (usable_w - 2 * img_gap) / 3  # width per image slot

    # Header / footer heights
    header_h = 18
    footer_h = 12
    img_area_h = page_h - 2 * margin - header_h - footer_h

    labels = ["模板图", "实拍图", "差异图"] if cjk_font_path else ["Template", "Target", "Diff"]

    for idx, record in enumerate(records):
        pdf.add_page()

        # Resolve result.json
        rj = _resolve_path(record.result_json) if record.result_json else None
        result_data: Dict = {}
        if rj and rj.exists():
            try:
                result_data = json.loads(rj.read_text(encoding="utf-8"))
            except (json.JSONDecodeError, OSError):
                pass

        # ── Header ──
        pdf.set_font(font_name, "B", 14)
        pdf.set_xy(margin, margin)
        pdf.cell(usable_w / 2, 8, f"#{idx + 1}  Code: {record.code}", new_x="RIGHT")

        pdf.set_font(font_name, "", 10)
        pdf.cell(usable_w / 2, 8, f"Case: {record.case_id}", new_x="LMARGIN", new_y="NEXT", align="R")

        # Verdict line
        verdict = record.verdict or result_data.get("verdict") or ""
        verdict = _strip_emoji(verdict)
        pdf.set_font(font_name, "B", 11)
        pdf.set_x(margin)
        pdf.cell(usable_w, 7, verdict, new_x="LMARGIN", new_y="NEXT")

        img_top = margin + header_h

        # ── Three images side by side ──
        images = [
            _resolve_image_for_record(record, rj, "template"),
            _resolve_image_for_record(record, rj, "target"),
            _resolve_image_for_record(record, rj, "diff"),
        ]

        for col, (img_path, label) in enumerate(zip(images, labels)):
            x = margin + col * (img_w + img_gap)

            # Label above image
            pdf.set_font(font_name, "B", 10)
            pdf.set_xy(x, img_top)
            pdf.cell(img_w, 5, label, align="C")

            img_y = img_top + 6

            if img_path and img_path.exists():
                try:
                    pdf.image(
                        str(img_path),
                        x=x,
                        y=img_y,
                        w=img_w,
                        h=0,  # auto height keeping aspect ratio
                    )
                except Exception:
                    pdf.set_xy(x, img_y + img_area_h / 2 - 3)
                    pdf.set_font(font_name, "", 9)
                    pdf.cell(img_w, 6, "(load error)", align="C")
            else:
                pdf.set_xy(x, img_y + img_area_h / 2 - 3)
                pdf.set_font(font_name, "", 9)
                pdf.cell(img_w, 6, "(no image)", align="C")

    report_path = Path(output_dir) / PDF_REPORT_FILENAME
    pdf.output(str(report_path))
    return report_path
